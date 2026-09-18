from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "generated"
OUTPUT_DIR.mkdir(exist_ok=True)


def set_cell_shading(cell, fill):
    """Set background color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def set_document_margins(section):
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)


def add_section_heading(document, text):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)

    # Add bottom border
    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")

    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")

    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0

    run = paragraph.add_run(text)
    run.font.size = Pt(8.5)

    return paragraph


def add_experience(document, experience):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)

    company = paragraph.add_run(experience.get("company", ""))
    company.bold = True
    company.font.size = Pt(9)

    paragraph.add_run(" | ").font.size = Pt(9)

    title = paragraph.add_run(experience.get("title", ""))
    title.italic = True
    title.font.size = Pt(9)

    dates = experience.get("dates", "")

    if dates:
        paragraph.add_run(f" | {dates}").font.size = Pt(8.5)

    for bullet in experience.get("bullets", []):
        add_bullet(document, bullet)


def add_project(document, project):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(1)

    name = paragraph.add_run(project.get("name", ""))
    name.bold = True
    name.font.size = Pt(8.8)

    for bullet in project.get("bullets", []):
        add_bullet(document, bullet)


def generate_resume_docx(
    tailored_resume: dict,
    candidate_name: str = "Stuti Bimali",
):
    """
    Generate a DOCX resume from the structured TailoredResume output.

    This function only formats information supplied by the pipeline.
    It does not create new resume facts.
    """

    document = Document()

    # Page setup
    section = document.sections[0]
    set_document_margins(section)

    # Global font
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(8)

    # -------------------------
    # HEADER
    # -------------------------

    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(1)

    name_run = header.add_run(candidate_name)
    name_run.bold = True
    name_run.font.size = Pt(17)

    target_title = tailored_resume.get("target_title", "")

    if target_title:
        title_paragraph = document.add_paragraph()
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_paragraph.paragraph_format.space_after = Pt(2)

        title_run = title_paragraph.add_run(target_title)
        title_run.font.size = Pt(9.5)
        title_run.bold = True

    # -------------------------
    # SUMMARY
    # -------------------------

    summary = tailored_resume.get("summary", "")

    if summary:
        add_section_heading(document, "Professional Summary")

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(summary)
        run.font.size = Pt(8.5)

    # -------------------------
    # SKILLS
    # -------------------------

    skills = tailored_resume.get("targeted_skills", [])

    if skills:
        add_section_heading(document, "Technical Skills")

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0

        run = paragraph.add_run(" | ".join(skills))
        run.font.size = Pt(8.5)

    # -------------------------
    # EXPERIENCE
    # -------------------------

    experiences = tailored_resume.get("experience", [])

    if experiences:
        add_section_heading(document, "Professional Experience")

        for experience in experiences:
            add_experience(document, experience)

    # -------------------------
    # PROJECTS
    # -------------------------

    projects = tailored_resume.get("projects", [])

    if projects:
        add_section_heading(document, "Projects")

        for project in projects:
            add_project(document, project)

    # -------------------------
    # AWARDS & CERTIFICATIONS
    # -------------------------

    awards_certifications = tailored_resume.get(
        "awards_certifications",
        []
    )

    if awards_certifications:
        add_section_heading(
            document,
            "Awards & Certifications"
        )

        for item in awards_certifications:
            if isinstance(item, dict):
                name = item.get("name", "")
            else:
                name = str(item)

            if name:
                add_bullet(document, name)
                
    # -------------------------
    # Save
    # -------------------------

    output_path = OUTPUT_DIR / "tailored_resume.docx"

    document.save(output_path)

    return output_path

if __name__ == "__main__":
    test_resume = {
        "target_title": "AI Engineer",
        "summary": (
            "AI/ML Engineer experienced in building LLM applications, "
            "RAG systems, AI agents and production APIs."
        ),
        "targeted_skills": [
            "Python",
            "LangChain",
            "LangGraph",
            "RAG",
            "FastAPI",
            "AWS",
            "FAISS"
        ],
        "experience": [
            {
                "company": "Example Company",
                "title": "AI Engineer",
                "dates": "2026",
                "bullets": [
                    "Built AI-powered applications using Python and LLM APIs.",
                    "Developed retrieval-based workflows for document processing."
                ]
            }
        ],
        "projects": [
            {
                "name": "AI Resume Tailor",
                "bullets": [
                    "Built an evidence-grounded resume tailoring workflow using LangChain and LangGraph."
                ]
            }
        ]
    }

    path = generate_resume_docx(test_resume)

    print(f"Resume generated: {path}")